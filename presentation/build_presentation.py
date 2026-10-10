from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "Personal_Finance_Manager_Review.pptx"

W = 13.333
H = 7.5

INK = "1D302A"
MUTED = "66756E"
FOREST = "183F35"
DEEP = "103229"
PAPER = "F2F5F0"
WHITE = "FFFFFF"
LINE = "D9E1DB"
CORAL = "D9634C"
CORAL_PALE = "FBEAE5"
CITRON = "DCE779"
MINT = "E1EEE7"
SAGE = "B8CFC2"
BLUE = "3B91CF"


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, radius=False, line=None, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=16, color=INK, font="Aptos", bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0, italic=False):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.3):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def arrow(slide, x, y, w=0.42, h=0.3, color=CORAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def circle(slide, x, y, diameter, fill, line_color=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(
            y), Inches(diameter), Inches(diameter)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line_color:
        shape.line.color.rgb = rgb(line_color)
    else:
        shape.line.fill.background()
    return shape


def image_fit(slide, path, x, y, w, h, background=WHITE):
    rect(slide, x, y, w, h, background, radius=True, line=LINE)
    from PIL import Image

    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = w / h
    if ratio >= box_ratio:
        image_w = w
        image_h = w / ratio
    else:
        image_h = h
        image_w = h * ratio
    left = x + (w - image_w) / 2
    top = y + (h - image_h) / 2
    slide.shapes.add_picture(str(path), Inches(
        left), Inches(top), Inches(image_w), Inches(image_h))


def background(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def footer(slide, number, dark=False):
    color = "B8CFC2" if dark else MUTED
    line(slide, 0.7, 7.08, 12.62, 7.08, "527267" if dark else LINE, 0.8)
    text(slide, "PERSONAL FINANCE MANAGER  /  PROJECT REVIEW", 0.72, 7.17, 7.7, 0.16,
         8, color, bold=True)
    text(slide, f"02 OCT 2026  ·  {number:02d}", 10.4, 7.17, 2.2, 0.16,
         8, color, align=PP_ALIGN.RIGHT)


def heading(slide, kicker, title, subtitle=None, number=None):
    text(slide, kicker.upper(), 0.76, 0.47, 8.5, 0.22, 9, CORAL, bold=True)
    text(slide, title, 0.74, 0.82, 11.9, 0.58,
         27, INK, font="Georgia", bold=True)
    if subtitle:
        text(slide, subtitle, 0.76, 1.48, 11.7, 0.45, 12, MUTED)
    if number:
        footer(slide, number)


def card(slide, x, y, w, h, title, body, accent=FOREST, fill=WHITE, index=None):
    rect(slide, x, y, w, h, fill, radius=True, line=LINE)
    rect(slide, x, y, 0.07, h, accent)
    if index:
        text(slide, index, x + 0.24, y + 0.22,
             w - 0.48, 0.2, 8, MUTED, bold=True)
        title_y = y + 0.56
    else:
        title_y = y + 0.28
    text(slide, title, x + 0.24, title_y, w - 0.48,
         0.4, 17, INK, font="Georgia", bold=True)
    text(slide, body, x + 0.24, title_y + 0.48, w - 0.48, h - (title_y - y) - 0.62,
         11, MUTED)


def add_note(slide, label, value, x, y, w, accent=CORAL):
    circle(slide, x, y + 0.03, 0.12, accent)
    text(slide, label, x + 0.24, y, w - 0.24, 0.25, 12, INK, bold=True)
    text(slide, value, x + 0.24, y + 0.3, w - 0.24, 0.62, 10, MUTED)


def new_slide(prs, color=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, color)
    return slide


def add_entity(slide, x, y, w, h, name, fields, accent=FOREST):
    rect(slide, x, y, w, h, WHITE, radius=True, line=LINE)
    rect(slide, x, y, w, 0.43, accent, radius=True)
    text(slide, name, x + 0.16, y + 0.08, w - 0.32, 0.25, 13, WHITE, bold=True)
    top = y + 0.58
    for index, field in enumerate(fields):
        color = CORAL if "PK" in field or "FK" in field else MUTED
        text(slide, field, x + 0.18, top + index * 0.29, w - 0.34, 0.22, 9, color,
             bold="PK" in field or "FK" in field)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Personal Finance Manager — Project Review"
    prs.core_properties.subject = "Java web application project presentation"
    prs.core_properties.author = "Personal Finance Manager team"
    prs.core_properties.keywords = "Java, Tomcat, MySQL, finance, project review"

    dashboard = ASSETS / "dashboard_welcome.png"
    dashboard_metrics = ASSETS / "dashboard_metrics.png"
    dashboard_actions = ASSETS / "dashboard_actions.png"
    admin_capture = ASSETS / "admin_create_capture.png"

    slide = new_slide(prs, DEEP)
    rect(slide, 0, 0, 0.16, H, CORAL)
    circle(slide, 10.7, -0.8, 3.5, FOREST)
    circle(slide, 11.8, 5.4, 1.0, CITRON)
    text(slide, "PROJECT REVIEW  /  JAVA WEB APPLICATION", 0.88, 0.72, 6.5, 0.25,
         10, CITRON, bold=True)
    text(slide, "Personal Finance\nManager", 0.84, 1.48, 7.2, 2.05, 36, WHITE,
         font="Georgia", bold=True)
    text(slide, "A practical workspace for tracking expenses, planning budgets, and keeping savings goals visible.",
         0.9, 3.86, 6.2, 0.9, 16, "D6E4DD")
    rect(slide, 0.9, 5.28, 5.85, 0.48, "244E42", radius=True)
    text(slide, "JAVA 11+   ·   SERVLETS + JSP   ·   TOMCAT 9   ·   MYSQL 8",
         1.06, 5.4, 5.55, 0.18, 9, WHITE, bold=True)
    text(slide, "REVIEW DECK  ·  02 OCTOBER 2026", 0.9, 6.55, 5.6, 0.25,
         10, "B8CFC2", bold=True)
    footer(slide, 1, dark=True)

    slide = new_slide(prs)
    heading(slide, "01 / PRODUCT", "Everyday money, in one focused workspace",
            "Three connected jobs keep personal finance activity organized.", 2)
    card(slide, 0.78, 2.15, 3.75, 2.35, "Track spending",
         "Record dated expenses with a category, amount, and optional description. Review or remove entries from your own list.",
         CORAL, WHITE, "01   /   EXPENSES")
    card(slide, 4.78, 2.15, 3.75, 2.35, "Plan with budgets",
         "Set a category limit for a monthly or yearly period, then keep budget totals visible alongside spending.",
         FOREST, MINT, "02   /   BUDGETS")
    card(slide, 8.78, 2.15, 3.75, 2.35, "Save toward goals",
         "Give a savings target a name, amount, and deadline so future priorities have a place in the plan.",
         CORAL, CORAL_PALE, "03   /   GOALS")
    text(slide, "CAPTURE", 0.78, 5.35, 3.75, 0.25, 10, FOREST,
         bold=True, align=PP_ALIGN.CENTER)
    text(slide, "PLAN", 4.78, 5.35, 3.75, 0.25, 10, FOREST,
         bold=True, align=PP_ALIGN.CENTER)
    text(slide, "REVIEW", 8.78, 5.35, 3.75, 0.25, 10, FOREST,
         bold=True, align=PP_ALIGN.CENTER)
    arrow(slide, 4.55, 5.3, color=CORAL)
    arrow(slide, 8.55, 5.3, color=CORAL)
    rect(slide, 0.82, 5.88, 11.7, 0.56, FOREST, radius=True)
    text(slide, "One account  →  personal records  →  a dashboard that points to the next useful action",
         1.05, 6.05, 11.15, 0.22, 12, WHITE, bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide(prs)
    heading(slide, "02 / EXPERIENCE", "Dashboard, from welcome to next action",
            "A composed view of three genuine dashboard regions from the running app.", 3)
    image_fit(slide, dashboard, 0.78, 2.04, 7.3, 3.48, WHITE)
    slide.shapes.add_picture(str(dashboard_metrics), Inches(0.78), Inches(5.68),
                             width=Inches(7.3), height=Inches(1.14))
    text(slide, "QUICK ACTIONS", 8.48, 2.08, 3.8, 0.24, 10, CORAL, bold=True)
    image_fit(slide, dashboard_actions, 8.48, 2.4, 3.8, 2.05, WHITE)
    add_note(slide, "Visible role", "The signed-in account and USER/ADMIN role remain easy to spot.",
             8.48, 4.72, 3.8, FOREST)
    add_note(slide, "Linked totals", "Expenses, budgets, and goals open their detail lists.",
             8.48, 5.62, 3.8, BLUE)
    text(slide, "KPI VALUES USE TEMPORARY REVIEW DATA", 0.8, 6.88, 7.2, 0.12,
         7, MUTED, bold=True)

    slide = new_slide(prs)
    heading(slide, "03 / DAILY WORKFLOWS", "The detail pages make the records useful",
            "Forms and lists keep entry simple and make saved data easy to review.", 4)
    workflow_shots = [
        (0.78, ASSETS / "expenses_table.png", "Expenses",
         "Date · category · amount · note"),
        (4.82, ASSETS / "budgets_table.png", "Budgets",
            "Period limits and stored spend"),
        (8.86, ASSETS / "goals_table.png", "Goals",
            "Targets · progress · deadline"),
    ]
    for x, image_path, title, caption in workflow_shots:
        image_fit(slide, image_path, x, 2.12, 3.68, 2.42, WHITE)
        text(slide, title, x + 0.08, 4.78, 3.5, 0.32, 17, INK,
             font="Georgia", bold=True)
        text(slide, caption, x + 0.08, 5.23, 3.5, 0.55, 10, MUTED)
    rect(slide, 0.84, 6.05, 11.65, 0.54, MINT, radius=True)
    text(slide, "The current user sees only their own records; admin reporting is a separate, role-guarded route.",
         1.05, 6.22, 11.2, 0.18, 10, FOREST, bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide(prs)
    heading(slide, "04 / ACCESS MODEL", "Registration is open; admin creation is not",
            "Role is assigned by the server and checked again on protected routes.", 5)
    rect(slide, 0.82, 2.04, 3.55, 0.72, WHITE, radius=True, line=LINE)
    text(slide, "PUBLIC REGISTRATION", 1.02, 2.29, 3.15, 0.2, 12, INK, bold=True,
         align=PP_ALIGN.CENTER)
    arrow(slide, 4.62, 2.24, 0.5, 0.3, CORAL)
    rect(slide, 5.28, 2.04, 3.15, 0.72, MINT, radius=True, line=LINE)
    text(slide, "USER ACCOUNT", 5.52, 2.29, 2.7, 0.2, 12, FOREST, bold=True,
         align=PP_ALIGN.CENTER)
    text(slide, "Registration always assigns USER", 8.78, 2.29, 3.6, 0.2,
         11, MUTED, italic=True)
    rect(slide, 0.82, 3.18, 3.55, 0.72, WHITE, radius=True, line=LINE)
    text(slide, "SHARED LOGIN", 1.02, 3.43, 3.15, 0.2, 12, INK, bold=True,
         align=PP_ALIGN.CENTER)
    arrow(slide, 4.62, 3.39, 0.5, 0.3, CORAL)
    rect(slide, 5.28, 3.05, 3.15, 0.72, MINT, radius=True, line=LINE)
    text(slide, "USER → DASHBOARD", 5.45, 3.3, 2.8, 0.2, 11, FOREST, bold=True,
         align=PP_ALIGN.CENTER)
    rect(slide, 5.28, 3.92, 3.15, 0.72, CORAL_PALE, radius=True, line=LINE)
    text(slide, "ADMIN → ADMIN AREA", 5.45, 4.17, 2.8, 0.2, 11, CORAL, bold=True,
         align=PP_ALIGN.CENTER)
    text(slide, "role-aware redirect", 8.78,
         3.5, 3.0, 0.2, 11, MUTED, italic=True)
    rect(slide, 0.82, 5.1, 5.45, 1.15, FOREST, radius=True)
    text(slide, "ADMIN-ONLY /admin", 1.08,
         5.34, 4.9, 0.24, 14, WHITE, bold=True)
    text(slide, "Create admin accounts  ·  inspect cross-user records  ·  reject non-admin sessions",
         1.08, 5.72, 4.9, 0.38, 10, "D6E4DD")
    rect(slide, 6.62, 5.1, 5.85, 1.15, CITRON, radius=True)
    text(slide, "USER AREA", 6.9, 5.34, 5.1, 0.24, 14, DEEP, bold=True)
    text(slide, "Manage personal expenses, budgets, goals, and password",
         6.9, 5.72, 5.1, 0.28, 10, DEEP)

    slide = new_slide(prs)
    heading(slide, "05 / ADMIN AREA", "Admin tools are protected, not public",
            "The admin route requires an authenticated ADMIN session before it can show the account form or records.", 6)
    image_fit(slide, admin_capture, 0.78, 2.12, 6.0, 3.55, WHITE)
    add_note(slide, "Create admin accounts", "Only an existing administrator can create another admin login.",
             7.35, 2.22, 4.9, CORAL)
    add_note(slide, "Review all records", "The admin view lists users and cross-user finance records read-only.",
             7.35, 3.54, 4.9, FOREST)
    add_note(slide, "Server-side guard", "Unauthenticated visitors go to login; non-admin users receive 403.",
             7.35, 4.86, 4.9, BLUE)
    text(slide, "SCREENSHOT CROPPED TO ACCOUNT-CREATION PANEL; USER EMAILS OMITTED",
         0.82, 5.92, 5.9, 0.22, 8, MUTED, bold=True)

    slide = new_slide(prs)
    heading(slide, "06 / DATA MODEL", "One account anchors three record types",
            "Every finance record carries a user_id foreign key; personal queries filter by the signed-in owner.", 7)
    line(slide, 3.72, 4.32, 5.55, 4.32, LINE, 2)
    line(slide, 5.55, 2.75, 5.55, 6.28, LINE, 2)
    line(slide, 5.55, 2.78, 7.35, 2.78, CORAL, 2)
    line(slide, 5.55, 4.52, 7.35, 4.52, FOREST, 2)
    line(slide, 5.55, 6.25, 7.35, 6.25, BLUE, 2)
    add_entity(slide, 0.82, 2.68, 2.9, 3.35, "app_user", [
        "id   PK", "username   UNIQUE", "email   UNIQUE", "password_hash", "role", "created_at"
    ], FOREST)
    add_entity(slide, 7.35, 2.08, 4.7, 1.4, "expense", [
        "id   PK", "user_id   FK", "date · category · amount · description"
    ], CORAL)
    add_entity(slide, 7.35, 3.82, 4.7, 1.4, "budget", [
        "id   PK", "user_id   FK", "category · period · limit · spent snapshot"
    ], FOREST)
    add_entity(slide, 7.35, 5.55, 4.7, 1.4, "goal", [
        "id   PK", "user_id   FK", "name · target · saved · deadline"
    ], BLUE)
    text(slide, "1 : many", 5.86, 2.58, 1.1, 0.2, 9, MUTED, bold=True)
    text(slide, "1 : many", 5.86, 4.32, 1.1, 0.2, 9, MUTED, bold=True)
    text(slide, "1 : many", 5.86, 6.05, 1.1, 0.2, 9, MUTED, bold=True)
    text(slide, "MySQL · InnoDB · foreign keys · unique usernames and emails",
         0.9, 6.35, 3.9, 0.38, 9, MUTED, bold=True)

    slide = new_slide(prs)
    heading(slide, "07 / ARCHITECTURE", "A compact MVC-style Java web stack",
            "The WAR packages application code, JSP views, and runtime dependencies for Tomcat.", 8)
    blocks = [
        (0.78, 2.65, 2.05, 1.62, "BROWSER", "JSP pages\nHTML forms", FOREST),
        (3.2, 2.65, 2.2, 1.62, "TOMCAT 9", "Servlet routes\nSession checks", CORAL),
        (5.78, 2.65, 2.0, 1.62, "SERVICES", "Auth rules\nBusiness logic", BLUE),
        (8.16, 2.65, 2.0, 1.62, "REPOSITORIES",
         "JDBC queries\nUser-scoped data", FOREST),
        (10.55, 2.65, 2.0, 1.62, "MYSQL 8",
         "app_user\nexpense · budget · goal", CORAL),
    ]
    for x, y, w, h, title, body, accent in blocks:
        rect(slide, x, y, w, h, WHITE, radius=True, line=LINE)
        rect(slide, x, y, w, 0.12, accent)
        text(slide, title, x + 0.15, y + 0.35,
             w - 0.3, 0.25, 11, accent, bold=True)
        text(slide, body, x + 0.15, y + 0.78,
             w - 0.3, 0.55, 11, INK, bold=True)
    for x in (2.83, 5.4, 7.78, 10.16):
        arrow(slide, x, 3.3, 0.3, 0.24, CORAL)
    rect(slide, 1.0, 5.18, 11.3, 0.82, "E8EEE9", radius=True)
    text(slide, "Java 11+   ·   Maven WAR   ·   javax.servlet 4   ·   JSP/JSTL   ·   JDBC + HikariCP",
         1.22, 5.43, 10.85, 0.25, 12, FOREST, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "Login dispatches by role; controllers forward to JSPs under /views.",
         1.2, 6.2, 10.9, 0.3, 11, MUTED, align=PP_ALIGN.CENTER)

    slide = new_slide(prs)
    heading(slide, "08 / TRUST & SECURITY", "The important boundaries are explicit",
            "Security comes from server-side checks and data-scoped queries—not hidden UI controls.", 9)
    card(slide, 0.82, 2.0, 5.65, 1.72, "Role checks on protected routes",
         "AdminServlet checks the signed-in session role before showing data or creating admins.",
         CORAL, WHITE, "AUTHORIZATION")
    card(slide, 6.82, 2.0, 5.65, 1.72, "BCrypt password storage",
         "PasswordUtil hashes passwords; password change checks the current password and requires 8+ characters.",
         FOREST, MINT, "CREDENTIALS")
    card(slide, 0.82, 3.94, 5.65, 1.72, "Owner-scoped operations",
         "Expense lookup, update, and delete repository methods include both record ID and user ID.",
         BLUE, WHITE, "DATA ISOLATION")
    card(slide, 6.82, 3.94, 5.65, 1.72, "Environment-based DB config",
         "FINANCE_DB_URL, FINANCE_DB_USER, and FINANCE_DB_PASSWORD stay outside source code.",
         CORAL, CORAL_PALE, "OPERATIONS")
    rect(slide, 0.84, 5.88, 11.6, 0.66, DEEP, radius=True)
    text(slide, "DATA NOTE  /  budget.spent_amount is stored and incremented when an expense is added.",
         1.08, 5.99, 11.12, 0.18, 8, CITRON, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "Expense edits/deletes are not reconciled yet, so this value can drift; add recomputation and tests next.",
         1.08, 6.24, 11.12, 0.18, 8, WHITE, align=PP_ALIGN.CENTER)

    slide = new_slide(prs)
    heading(slide, "09 / RUN LOCALLY", "Build the WAR, deploy to Tomcat, open the app",
            "The repository includes setup instructions in README.md and a manual evaluator walkthrough in TESTING.md.", 10)
    steps = [
        (0.84, "01", "Prepare", "Java 11+ · Maven\nMySQL 8 · Tomcat 9"),
        (3.92, "02", "Create schema",
         "Run src/main/resources/db/schema.sql\nSet FINANCE_DB_* for Tomcat"),
        (7.0, "03", "Package", "mvn clean package\nBuild target/personal-finance-web.war"),
        (10.08, "04", "Deploy", "Copy WAR to Tomcat webapps\nOpen /personal-finance-web/"),
    ]
    for x, number, title, body in steps:
        circle(slide, x + 0.07, 2.36, 0.48,
               CORAL if number in ("01", "03") else FOREST)
        text(slide, number, x + 0.07, 2.5, 0.48, 0.15, 9, WHITE, bold=True,
             align=PP_ALIGN.CENTER)
        if x < 10:
            line(slide, x + 0.55, 2.6, x + 2.7, 2.6, LINE, 2)
        text(slide, title, x, 3.08, 2.55, 0.34,
             17, INK, font="Georgia", bold=True)
        text(slide, body, x, 3.56, 2.6, 0.72, 10, MUTED)
    rect(slide, 0.84, 5.05, 11.65, 1.12, "FFF4D6", radius=True, line="E8D8A9")
    text(slide, "DATABASE SETUP", 1.08, 5.28, 2.1, 0.22, 9, CORAL, bold=True)
    text(slide, "Schema seeds demo roles for local review.", 3.2, 5.24, 5.3, 0.32,
         12, INK, bold=True)
    text(slide, "Change demo passwords before public use; keep real DB secrets in environment variables.",
         1.08, 5.69, 10.95, 0.22, 9, MUTED, italic=True)

    slide = new_slide(prs)
    heading(slide, "10 / REVIEW EVIDENCE", "A clear walkthrough; coverage can keep growing",
            "Seven focused unit tests now cover password hashing and authentication behavior.", 11)
    rect(slide, 0.82, 2.05, 4.05, 3.95, FOREST, radius=True)
    text(slide, "BUILD CHECK", 1.12, 2.42, 3.4, 0.25, 10, CITRON, bold=True)
    text(slide, "PASS", 1.12, 2.9, 3.4, 0.65,
         36, WHITE, font="Georgia", bold=True)
    text(slide, "mvn clean package\n\nWAR artifact: target/personal-finance-web.war",
         1.12, 3.86, 3.25, 0.95, 13, "D6E4DD")
    text(slide, "Build verified for this review deck on 02 Oct 2026.",
         1.12, 5.35, 3.28, 0.38, 9, "B8CFC2")
    add_note(slide, "Create a regular account", "Public registration assigns USER; log in and review the dashboard.",
             5.45, 2.18, 6.7, FOREST)
    add_note(slide, "Exercise the core workflows", "Add an expense, budget, and goal; verify lists and dashboard counts.",
             5.45, 3.35, 6.7, CORAL)
    add_note(slide, "Check access boundaries", "After logout, request protected routes; try cross-user record IDs.",
             5.45, 4.52, 6.7, BLUE)
    text(slide, "Next: repository/DB integration tests and data-isolation coverage. The manual walkthrough remains in TESTING.md.",
         5.7, 5.9, 6.35, 0.55, 10, MUTED, italic=True)

    slide = new_slide(prs)
    heading(slide, "11 / PROJECT TEAM", "Planned ownership across the team",
            "A proposed division of responsibility; update the slide as each contribution is completed.", 12)
    team = [
        (0.82, "SB", "Sachin Bhati", "JAVA BACKEND & SECURITY",
         "Authentication, role guards, password changes, and service behavior.",
         "STRETCH: reconcile budgets after expense edits/deletes; add regression tests.", FOREST, MINT),
        (4.82, "AB", "Aayush Bagauli", "DATA & INTEGRATION",
         "MySQL schema, JDBC persistence and owner scoping; Tomcat/Maven setup and verification.",
         "Own the end-to-end checklist, setup notes, and review presentation.", CORAL, CORAL_PALE),
        (8.82, "AP", "Ayush Pandey", "FRONTEND & UX",
         "JSP/CSS screens, responsive forms and tables, dashboard workflows, and accessibility checks.",
         "Capture and review the primary user and admin flows at desktop and mobile sizes.", BLUE, "EAF1F5"),
    ]
    for x, initials, name, area, description, stretch, accent, tint in team:
        rect(slide, x, 2.12, 3.7, 3.76, WHITE, radius=True, line=LINE)
        rect(slide, x, 2.12, 3.7, 0.1, accent)
        circle(slide, x + 0.26, 2.43, 0.58, tint)
        text(slide, initials, x + 0.26, 2.62, 0.58, 0.16, 10, accent, bold=True,
             align=PP_ALIGN.CENTER)
        text(slide, name, x + 0.98, 2.53, 2.45, 0.32,
             14, INK, font="Georgia", bold=True)
        text(slide, area, x + 0.28, 3.22, 3.12, 0.24, 9, accent, bold=True)
        text(slide, description, x + 0.28, 3.68, 3.12, 0.82, 10, MUTED)
        line(slide, x + 0.28, 4.66, x + 3.42, 4.66, LINE, 0.8)
        text(slide, stretch, x + 0.28, 4.87, 3.12, 0.72, 9, INK, bold=True)
    rect(slide, 0.84, 6.12, 11.65, 0.48, DEEP, radius=True)
    text(slide, "PLANNED RESPONSIBILITIES  ·  Confirm task handoffs as work is completed.",
         1.08, 6.28, 11.12, 0.16, 9, WHITE, bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide(prs, DEEP)
    rect(slide, 0, 0, 0.16, H, CORAL)
    text(slide, "TAKEAWAY", 0.9, 0.76, 3.0, 0.22, 10, CITRON, bold=True)
    text(slide, "A small finance app,\nwith clear next steps.", 0.86, 1.48, 8.5, 1.7,
         34, WHITE, font="Georgia", bold=True)
    text(slide, "Track everyday records. Keep personal data scoped. Give reviewers a predictable path through the product.",
         0.92, 3.55, 7.2, 0.72, 15, "D6E4DD")
    rect(slide, 8.8, 1.62, 3.45, 3.8, "244E42", radius=True)
    text(slide, "REVIEW PATH", 9.12, 1.98, 2.8, 0.2, 9, CITRON, bold=True)
    text(slide, "Register as USER\n↓\nLog a transaction\n↓\nPlan a budget or goal\n↓\nReview as ADMIN",
         9.12, 2.45, 2.8, 2.45, 16, WHITE, font="Georgia", bold=True)
    text(slide, "NEXT QUALITY STEP", 0.92, 5.6, 2.2, 0.22, 9, CORAL, bold=True)
    text(slide, "Expand coverage to repository/database integration tests, then repeat the documented data-isolation checks.",
         0.92, 5.98, 7.2, 0.48, 12, "B8CFC2")
    footer(slide, 13, dark=True)

    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_deck()
